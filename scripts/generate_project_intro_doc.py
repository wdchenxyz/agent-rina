#!/usr/bin/env python3
"""Generate the Agent Rina project introduction DOCX.

Run with the Codex bundled Python runtime; the script expects the diagram PNGs
under docs/project-introduction/diagrams/.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DIAGRAM_DIR = ROOT / "docs" / "project-introduction" / "diagrams"
OUTPUT = ROOT / "docs" / "Agent-Rina-Project-Guide.zh-TW.docx"

# Resolved design tokens: standard_business_brief + editorial_cover header pattern.
PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(1.0)
HEADER_DISTANCE = Inches(0.492)
FOOTER_DISTANCE = Inches(0.492)
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}

# Named typography override requested for the editable DOCX. Word selects
# these by script: Aptos for Latin text and PingFang TC for Traditional Chinese.
BODY_FONT = "Aptos"
EAST_ASIA_FONT = "PingFang TC"
INK = "263238"
MUTED = "5E6E79"
HEADING_BLUE = "2E74B5"
HEADING_DARK_BLUE = "1F4D78"
NAVY = "183A61"
GOLD = "B48116"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
PALE_GOLD = "FFF5D8"
PALE_TEAL = "E3F2EF"
PALE_RED = "FCE8E4"
WHITE = "FFFFFF"


def set_run_font(
    run,
    *,
    name: str = BODY_FONT,
    east_asia: str = EAST_ASIA_FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK, align=None, size=9.3):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    if align is not None:
        p.alignment = align
    set_run_font(p.add_run(text), size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width_dxa(cell, width: int):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:type"), "dxa")
    tcw.set(qn("w:w"), str(width))


def set_table_geometry(table, widths: list[int], *, indent: int = TABLE_INDENT_DXA):
    assert sum(widths) == CONTENT_WIDTH_DXA, (widths, sum(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(CONTENT_WIDTH_DXA))

    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:type"), "dxa")
    tblind.set(qn("w:w"), str(indent))

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tc_mar = tblpr.find(qn("w:tblCellMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tblCellMar")
        tblpr.append(tc_mar)
    for edge, val in CELL_MARGINS.items():
        el = tc_mar.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for idx, cell in enumerate(row.cells):
            set_cell_width_dxa(cell, widths[idx])


def set_table_borders(table, color="CAD2D9", size="6"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def add_field(paragraph, instruction: str):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def add_numbering_definition(doc: Document, *, bullet: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    numfmt = OxmlElement("w:numFmt")
    numfmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl.append(numfmt)
    lvltext = OxmlElement("w:lvlText")
    lvltext.set(qn("w:val"), "•" if bullet else "%1.")
    lvl.append(lvltext)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)
    if bullet:
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), "Arial")
        rfonts.set(qn("w:hAnsi"), "Arial")
        rpr.append(rfonts)
        lvl.append(rpr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, numid])


def add_bullet(doc, text: str, bullet_num_id: int, *, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="Normal")
    apply_numbering(p, bullet_num_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True, color=INK)
        set_run_font(p.add_run(text[len(bold_prefix):]), color=INK)
    else:
        set_run_font(p.add_run(text), color=INK)
    return p


def add_numbered(doc, text: str, number_num_id: int):
    p = doc.add_paragraph(style="Normal")
    apply_numbering(p, number_num_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    set_run_font(p.add_run(text), color=INK)
    return p


def add_body(doc, text: str, *, bold_lead: str | None = None, keep_with_next=False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.keep_with_next = keep_with_next
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True, color=INK)
        set_run_font(p.add_run(text[len(bold_lead):]), color=INK)
    else:
        set_run_font(p.add_run(text), color=INK)
    return p


def add_callout(doc, label: str, text: str, *, fill=PALE_GOLD, border=GOLD):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.keep_together = True
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border)
    borders.append(left)
    ppr.append(borders)
    set_run_font(p.add_run(f"{label}  "), bold=True, color=NAVY)
    set_run_font(p.add_run(text), color=INK)
    return p


def add_figure(doc, png_name: str, caption: str, alt_text: str):
    path = DIAGRAM_DIR / png_name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(6.5))
    docpr = shape._inline.docPr
    docpr.set("name", caption)
    docpr.set("descr", alt_text)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.keep_with_next = False
    set_run_font(cap.add_run(caption), size=9, color=MUTED, italic=True)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, HEADING_BLUE, 16, 8),
        "Heading 2": (13, HEADING_BLUE, 12, 6),
        "Heading 3": (12, HEADING_DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def configure_sections(doc: Document):
    for section in doc.sections:
        section.page_width = PAGE_WIDTH
        section.page_height = PAGE_HEIGHT
        section.top_margin = MARGIN
        section.right_margin = MARGIN
        section.bottom_margin = MARGIN
        section.left_margin = MARGIN
        section.header_distance = HEADER_DISTANCE
        section.footer_distance = FOOTER_DISTANCE
        section.different_first_page_header_footer = True

        header = section.header
        hp = header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        set_run_font(hp.add_run("AGENT RINA  •  PROJECT GUIDE"), size=8.5, color=MUTED, bold=True)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_before = Pt(0)
        set_run_font(fp.add_run("PROJECT SNAPSHOT  •  2026-07-25    "), size=8.5, color=MUTED)
        add_field(fp, "PAGE")


def build_document() -> Document:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN
    section.right_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.header_distance = HEADER_DISTANCE
    section.footer_distance = FOOTER_DISTANCE
    section.different_first_page_header_footer = True

    props = doc.core_properties
    props.title = "Agent Rina 專案介紹與架構導覽"
    props.subject = "Architecture and project guide for Agent Rina"
    props.author = "Codex"
    props.keywords = "Agent Rina, Next.js, Chat SDK, AI SDK, Slack, Telegram, architecture"
    props.comments = "Snapshot based on the repository state on 2026-07-25."

    bullet_num_id = add_numbering_definition(doc, bullet=True)
    number_num_id = add_numbering_definition(doc, bullet=False)
    lifecycle_num_id = add_numbering_definition(doc, bullet=False)
    roadmap_num_id = add_numbering_definition(doc, bullet=False)

    # Cover — editorial_cover named override.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(84)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("PROJECT & ARCHITECTURE GUIDE"), size=11, color=GOLD, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run("Agent Rina"), size=34, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    set_run_font(p.add_run("多平台 AI 助理的專案介紹與架構導覽"), size=17, color=HEADING_DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.55)
    p.paragraph_format.right_indent = Inches(0.55)
    p.paragraph_format.space_after = Pt(48)
    p.paragraph_format.line_spacing = 1.25
    set_run_font(
        p.add_run("從 Slack／Telegram 介接、對話生命週期與 ToolLoopAgent，\n一路看到工具路由、檔案交付與安全邊界。"),
        size=12.5,
        color=MUTED,
    )

    # Hand-drawn motif.
    motif = doc.add_paragraph()
    motif.alignment = WD_ALIGN_PARAGRAPH.CENTER
    motif.paragraph_format.space_after = Pt(52)
    set_run_font(motif.add_run("Slack  →  Webhook  →  Rina Core  ⇄  Tools  →  Reply"), size=13, color=NAVY, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(4)
    set_run_font(meta.add_run("Repository snapshot  •  2026-07-25"), size=10.5, color=MUTED, bold=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("Next.js 16  •  Chat SDK 4.26  •  AI SDK 6  •  TypeScript"), size=9.5, color=MUTED)

    add_page_break(doc)

    # 0. Reader orientation.
    doc.add_heading("閱讀導覽", level=1)
    add_body(
        doc,
        "這份文件是面向新加入的工程師與專案利害關係人的技術導覽。它不是逐行 API 文件，而是說明 Agent Rina 的定位、主要元件、訊息如何流動、外部工具在哪裡被呼叫，以及目前有哪些可見的限制。",
    )
    add_callout(
        doc,
        "一句話定位",
        "Agent Rina 是一個以同一套 AI 核心服務 Slack 與 Telegram 的多平台聊天機器人；平台差異留在介接與輸出層，推理、工具選擇與安全控制集中在共享的 Rina runtime。",
        fill=PALE_GOLD,
    )

    doc.add_heading("專案快照", level=2)
    for text, prefix in [
        ("應用框架：Next.js 16 App Router；Webhook 入口位於動態平台路由。", "應用框架："),
        ("對話介接：Chat SDK 搭配 Slack、Telegram adapter 與 Redis state adapter。", "對話介接："),
        ("AI 執行：AI SDK ToolLoopAgent，透過 AI Gateway 呼叫模型，單次最多 20 steps。", "AI 執行："),
        ("工具能力：Web research、網頁內容與圖片擷取、arXiv 原始碼、Python sandbox、artifact 下載／上傳、bash 與 repo-local skills。", "工具能力："),
        ("主要語言：TypeScript；套件管理以 pnpm 為主，並保留 npm lockfile。", "主要語言："),
    ]:
        add_bullet(doc, text, bullet_num_id, bold_prefix=prefix)

    doc.add_heading("設計主軸", level=2)
    add_body(doc, "目前實作最清楚的三個設計選擇如下：")
    for text in [
        "共享核心：handlers、agent、tool registry 與 delivery bridge 不綁定單一聊天平台。",
        "意圖先行：planner 先以關鍵字與媒體類型判斷請求類別，前兩個推理步驟只啟用較相關的工具。",
        "I/O 邊界明確：存取控制、附件上限、safe fetch、sandbox、artifact 路徑與大小限制分散在實際執行邊界，而不是只靠模型提示。",
    ]:
        add_numbered(doc, text, number_num_id)

    add_page_break(doc)

    # 1. System architecture.
    doc.add_heading("1. 系統架構全貌", level=1)
    add_body(
        doc,
        "整體可以拆成四個區域：使用者通道、Web 邊緣、Rina 共享核心，以及外部模型／狀態／工具服務。訊息由平台 Webhook 進入，經 Chat SDK 統一成 thread 與 message，再交由 Rina 處理；回覆則沿 delivery bridge 回到原平台。",
    )
    add_figure(
        doc,
        "system-context.png",
        "圖 1　Agent Rina 系統情境圖（tldraw 風格）",
        "Slack and Telegram enter a Next.js webhook and Chat SDK adapters, then pass through shared Rina handlers, an agent runtime, tools, and a platform-aware delivery bridge.",
    )
    add_callout(
        doc,
        "架構重點",
        "Slack 與 Telegram 並不是兩套 bot；它們共享同一個 agent 與工具層。平台差異主要集中在事件觸發、處理中指示與回覆投遞方式。",
        fill=PALE_TEAL,
        border="3C9B88",
    )

    add_page_break(doc)
    doc.add_heading("主要元件與責任", level=2)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1656, 3456, 4248])
    set_table_borders(table)
    headers = ["層次", "主要檔案", "責任"]
    for i, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)
        shade_cell(table.rows[0].cells[i], PALE_BLUE)
    repeat_table_header(table.rows[0])
    component_rows = [
        ("Web 入口", "app/api/webhooks/[platform]/route.ts", "依 platform 尋找 webhook handler；用 Next.js after() 承接背景工作並記錄失敗。"),
        ("平台介接", "lib/bot/index.ts", "依環境變數建立 Slack／Telegram adapter；使用 Redis state；啟動時註冊共享 handlers。"),
        ("生命週期", "lib/rina/handlers.ts", "Mention、訂閱訊息與 Telegram DM 的共同流程；存取檢查、訂閱、附件補抓、錯誤處理與狀態指示。"),
        ("模型輸入", "lib/rina/message-input.ts", "將文字、圖片、PDF、text file 與 thread history 轉成 AI SDK message；合併連續同角色訊息並限制歷史長度。"),
        ("推理核心", "lib/rina/agent.ts\nlib/rina/planner.ts", "建立 ToolLoopAgent、選模型、限制 steps、依 intent 縮小初始工具集，並把工具呼叫／結果接到 logger。"),
        ("工具層", "lib/rina/tools/*", "Web、網頁圖片、arXiv、sandbox、artifact 與結構化結果；registry 維護分類、狀態文案與 intent mapping。"),
        ("輸出層", "lib/rina/delivery.ts", "在 Slack 串流、Telegram 緩衝；顯示工具狀態、攔截 file-upload 結果、切分長訊息並重試暫時性錯誤。"),
    ]
    for ridx, row in enumerate(component_rows, start=1):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, bold=(i == 0), align=WD_ALIGN_PARAGRAPH.LEFT, size=8.9)
            if ridx % 2 == 0:
                shade_cell(cells[i], "FAFBFC")
    set_table_geometry(table, [1656, 3456, 4248])

    add_page_break(doc)

    # 2. Request lifecycle.
    doc.add_heading("2. 一次請求如何走完整條鏈路", level=1)
    add_body(
        doc,
        "下圖把一次請求濃縮成七個階段。真正的迴圈出現在 agent 與工具之間：模型可根據前一步工具結果繼續推理，直到產生回覆或達到 20-step 上限。",
    )
    add_figure(
        doc,
        "request-lifecycle.png",
        "圖 2　請求生命週期與 agent–tool 迴圈",
        "A seven-step request lifecycle: receive, gate and enrich, build context, plan intent, run ToolLoopAgent, execute tools, and deliver the response, with tool results looping back into reasoning.",
    )

    doc.add_heading("流程細節", level=2)
    lifecycle = [
        "事件觸發：Slack／Telegram mention 會建立或進入 thread；訂閱後的訊息會持續交給同一 handler。Telegram 私訊不需要 mention，另由 catch-all handler 接住。",
        "存取與補強：平台專屬 allowlist 優先於全域 allowlist。Slack app_mention 可能缺附件，因此沒有附件時會重新 fetch 原訊息。",
        "脈絡組裝：新訊息轉成 text、image 或 file parts；follow-up 另取得 thread history 與可能存在的新聞摘要 context，兩者平行讀取。",
        "意圖規劃：planner 判斷 paper、research、data、codebase、artifact、file 或 chat。這是工具相關性優化，不是永久權限隔離。",
        "推理與工具：前兩個 steps 只開啟 intent 對應的工具；之後 agent 可以看見完整工具集。所有工具結果再回到模型，直到形成回答。",
        "交付：Slack 使用串流橋接；Telegram 累積一段文字後再送，並把超過 4,000 字的內容分割。結構化 file-upload 結果會被攔截成聊天附件。",
    ]
    for text in lifecycle:
        add_numbered(doc, text, lifecycle_num_id)

    doc.add_heading("平台行為差異", level=2)
    add_body(doc, "Slack", bold_lead="Slack")
    for text in [
        "新 mention 會訂閱 thread，後續訊息由 onSubscribedMessage 處理。",
        "收到請求時最佳努力加上 eyes；成功後移除 eyes 並加上 check，失敗時清掉指示。",
        "文字以 AsyncIterable 即時串流；工具狀態與檔案另行張貼。",
    ]:
        add_bullet(doc, text, bullet_num_id)
    add_body(doc, "Telegram", bold_lead="Telegram")
    for text in [
        "群組／supergroup 依 mention 觸發；DM 由廣泛訊息 matcher 接住。",
        "文字緩衝後送出，長訊息優先在段落或換行處切分。",
        "若 Telegram 拒絕 Markdown entity，delivery 會改以 raw text 重送。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_page_break(doc)

    # 3. Tools and safety.
    doc.add_heading("3. 工具路由與安全邊界", level=1)
    add_body(
        doc,
        "Agent Rina 的工具設計分成兩層：registry 決定『什麼情境先看見哪些工具』，每一個真正接觸網路、檔案或執行環境的工具再自行落實安全限制。這兩層用途不同，不能互相取代。",
    )
    add_figure(
        doc,
        "tool-safety.png",
        "圖 3　意圖路由與外部 I/O 安全邊界",
        "The intent planner narrows tools for relevance, while safe fetch, isolated execution, artifact guards, access control, attachment limits, and delivery fallbacks enforce execution safety.",
    )

    doc.add_heading("工具地圖", level=2)
    tool_groups = [
        ("Web research：", "webSearch 透過 Gemini grounding 產生附來源回答；perplexitySearch 偏即時資訊；researchWeb 平行交叉兩者；fetchWebpage 讀單一 URL。"),
        ("網頁內容與圖片：", "extractWebpageAssets 以 safe fetch 取得 HTML，使用 Defuddle 擷取內容、從 metadata／JSON-LD／DOM 找圖、排序候選並把通過驗證的圖片轉成 file-upload 結果。"),
        ("研究論文：", "downloadArxivSource 下載並解壓 LaTeX source；listPaperFiles、readPaperFile 與 uploadPaperFigure 支援分段理解與圖表回傳。"),
        ("資料與執行：", "runPythonCode 在 Vercel Sandbox microVM 執行 Python，可安裝套件、擷取 stdout/stderr，並將指定輸出檔送回 artifacts。"),
        ("程式庫操作：", "bash-tool 與 skill tool 延遲載入；bash upload scope 限制在 app、lib、docs、.agents 內的程式與設定類檔案。"),
        ("Artifact：", "downloadFile、listArtifacts、uploadArtifact 管理 artifacts/，並透過結構化結果讓 delivery 統一投遞。"),
    ]
    for lead, body in tool_groups:
        add_body(doc, lead + body, bold_lead=lead)

    # Keep the safety matrix together after the font override; fallback
    # renderers otherwise strand its header row at the foot of the prior page.
    add_page_break(doc)
    doc.add_heading("具體防護", level=2)
    security_rows = [
        ("存取控制", "平台 user/chat allowlist 若有設定就優先使用；否則回退全域清單。user 與 chat 兩項同時設定時必須同時符合；全部未設時預設允許。"),
        ("輸入媒體", "最多送 4 個支援附件給模型；圖片上限 5 MB，PDF／text file 上限 10 MB。超出或無法讀取時跳過並提示。"),
        ("Safe fetch", "只允許 HTTP(S)，拒絕 embedded credentials、localhost／內部 hostname、私有／保留 IP；每次 redirect 重新驗證 DNS，預設 12 秒、最多 4 次 redirect，且有 content length 與串流 byte cap。"),
        ("Sandbox", "microVM 最長 2 分鐘、命令 90 秒；最多回收 8 個輸出檔，每檔上限 8 MB，執行結束後停止 sandbox。"),
        ("Artifacts / arXiv", "用 resolve + base prefix 防 path traversal；上傳副檔名白名單；單一聊天上傳 8 MB。arXiv 讀檔另有 100 KB 截斷。"),
        ("輸出與記錄", "網路暫時性錯誤採 400/1200/2500 ms 退避；Telegram Markdown 失敗回退 raw text；logger 對 base64 內容只記錄長度。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2700, 6660])
    set_table_borders(table)
    for i, text in enumerate(["邊界", "目前實作"]):
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)
        shade_cell(table.rows[0].cells[i], PALE_BLUE)
    repeat_table_header(table.rows[0])
    for ridx, row in enumerate(security_rows, start=1):
        cells = table.add_row().cells
        set_cell_text(cells[0], row[0], bold=True, size=9.1)
        set_cell_text(cells[1], row[1], size=8.9)
        if ridx % 2 == 0:
            shade_cell(cells[0], "FAFBFC")
            shade_cell(cells[1], "FAFBFC")
    set_table_geometry(table, [2700, 6660])

    add_page_break(doc)

    # 4. Data, state and observability.
    doc.add_heading("4. 對話狀態、附件與可觀測性", level=1)
    doc.add_heading("Thread state 與歷史", level=2)
    add_body(
        doc,
        "Chat SDK 的 state adapter 由 Redis 提供。follow-up 請求會讀取整個 thread，略過本次訊息，將 bot 訊息映射為 assistant、其他人訊息映射為 user；連續同角色訊息會合併，以符合部分 provider 對角色交替的要求。最後保留最近 RINA_MAX_HISTORY_MESSAGES 筆，預設為 30。",
    )
    add_body(
        doc,
        "Bot 自己張貼的附件會被轉成一則合成 user context：『This is the file you just posted to the chat.』如此一來，後續使用者說『剛才那張圖』時，模型仍能看見它。",
    )

    doc.add_heading("新聞摘要 context", level=2)
    add_body(
        doc,
        "lib/bot/news/ 目前已有 RSS 抓取與去重、AI 篩選摘要，以及把 digest 與 stories 存入 Redis 的 primitives。digest thread context 使用 news:digest-thread: 前綴，TTL 為 14 天；follow-up handler 會把找到的 digest 以前導訊息加入模型脈絡。",
    )
    add_callout(
        doc,
        "現況提醒",
        "目前程式樹沒有找到排程、發文或 saveDigestThreadContext 的實際呼叫點。instrumentation.ts 的註解寫著會 eager import bot 以啟動 node-cron，但 lib/bot/index.ts 尚未建立 scheduler。可把新聞功能視為『核心零件已存在、產品流程尚未串接』。",
        fill=PALE_RED,
        border="C75B4B",
    )

    doc.add_heading("記錄與除錯", level=2)
    for text in [
        "ThreadLogger 由 handler 以平台與 thread ID 建立，記錄分隔線、輸入、歷史、完整 prompt、plan、active tools、工具呼叫／結果與最終文字區塊。",
        "withLogging() 只是 tap stream，不改變下游看到的事件；工具輸出若包含 dataBase64，log 只留下字串長度。",
        "RINA_IMAGE_DEBUG=1 可開啟附件與模型 message 組裝細節，適合追查 Slack 附件遺失、MIME 推斷或歷史媒體問題。",
        "Handler boundary 捕捉整次執行錯誤，向使用者回覆簡短 internal error，並在 finally 內處理 Slack reaction。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    # 5. Configuration and operation.
    # Begin the configuration block on a fresh page so the environment matrix
    # stays intact across Aptos/PingFang and metric-compatible fallbacks.
    add_page_break(doc)
    doc.add_heading("5. 設定、開發與運行", level=1)
    doc.add_heading("必要與選用環境變數", level=2)
    env_rows = [
        ("AI_GATEWAY_API_KEY", "必要", "模型 Gateway 金鑰。"),
        ("REDIS_URL", "必要", "Chat SDK thread state 與 digest context。"),
        ("SLACK_BOT_TOKEN + SLACK_SIGNING_SECRET", "成對", "Slack adapter；只設定其中一個會在啟動時丟錯。"),
        ("TELEGRAM_BOT_TOKEN", "選用", "Telegram adapter；需公開 HTTPS webhook。"),
        ("TELEGRAM_WEBHOOK_SECRET", "建議", "Telegram webhook 驗證。"),
        ("BOT_USERNAME", "選用", "Bot 顯示名稱，預設 mybot。"),
        ("BOT_ALLOWED_*", "選用", "全域或分平台 user/chat allowlist。"),
        ("RINA_MAX_HISTORY_MESSAGES", "選用", "送入模型的 thread history 上限，預設 30。"),
        ("RINA_IMAGE_DEBUG", "選用", "設為 1 以記錄媒體組裝細節。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3456, 1584, 4320])
    set_table_borders(table)
    for i, text in enumerate(["變數", "狀態", "用途"]):
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)
        shade_cell(table.rows[0].cells[i], PALE_BLUE)
    repeat_table_header(table.rows[0])
    for ridx, row in enumerate(env_rows, start=1):
        cells = table.add_row().cells
        set_cell_text(cells[0], row[0], bold=True, size=8.6)
        set_cell_text(cells[1], row[1], align=WD_ALIGN_PARAGRAPH.CENTER, size=8.8)
        set_cell_text(cells[2], row[2], size=8.8)
        if ridx % 2 == 0:
            for cell in cells:
                shade_cell(cell, "FAFBFC")
    set_table_geometry(table, [3456, 1584, 4320])

    doc.add_heading("常用工作流程", level=2)
    commands = [
        ("安裝", "pnpm install"),
        ("開發伺服器", "pnpm dev  →  http://localhost:3000"),
        ("Slack webhook", "POST /api/webhooks/slack"),
        ("Telegram webhook", "POST /api/webhooks/telegram"),
        ("CLI playground", "pnpm pg \"your question\""),
        ("型別／Lint／Build", "pnpm exec tsc --noEmit  ·  pnpm lint  ·  pnpm build"),
    ]
    for label, command in commands:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(6)
        set_run_font(p.add_run(f"{label}　"), bold=True, color=NAVY)
        set_run_font(p.add_run(command), name="Menlo", east_asia=EAST_ASIA_FONT, size=9.6, color=INK)

    add_callout(
        doc,
        "啟動條件",
        "Slack 與 Telegram 至少要成功建立一個 adapter；若兩者都沒有配置，模組載入時直接失敗。instrumentation.ts 在 Node.js runtime 會 eager import bot，因此配置錯誤可能在應用啟動階段立即暴露。",
        fill=PALE_TEAL,
        border="3C9B88",
    )

    add_page_break(doc)

    # 6. Risks and extension points.
    doc.add_heading("6. 已知限制與下一步", level=1)
    add_body(
        doc,
        "以下不是抽象的最佳實務清單，而是依目前程式樹可以直接觀察到的工程風險與擴充點。",
    )
    risks = [
        ("新聞 digest 未完整串接", "完成 scheduler、目標 channel／thread 設定、實際發文與 saveDigestThreadContext 呼叫，並讓 instrumentation 註解與實作一致。"),
        ("Planner 主要依英文關鍵字", "中文或其他語言的研究、資料與 codebase 請求容易落到 chat intent；可加入多語詞彙、輕量分類模型或以模型做低成本 router。"),
        ("缺少自動化測試", "目前沒有看到測試檔；優先為 planner、allowlist fallback、safe-fetch URL 判斷、Telegram splitting 與 structured upload interception 建立單元測試。"),
        ("模型與供應商識別寫死", "主 agent 與新聞 curator 使用不同且硬編碼的 model ID；可移到環境設定並加入明確 fallback／成本與延遲策略。"),
        ("工具開放策略需清楚記錄", "activeTools 只在前兩 steps 收斂，之後回到完整 toolset；若未來需要真正 permission boundary，應把授權放到工具執行前，而非依賴 planner。"),
        ("依賴與 lockfile 需統一", "專案同時有 pnpm-lock.yaml 與 package-lock.json；可確立唯一套件管理流程，避免 CI 與本機解析出不同依賴。"),
    ]
    for title, recommendation in risks:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        set_run_font(p.add_run(title), size=11.2, color=HEADING_DARK_BLUE, bold=True)
        add_body(doc, recommendation)

    doc.add_heading("建議擴充順序", level=2)
    roadmap = [
        "先補測試與 configuration validation，建立安全重構的基線。",
        "完成新聞 digest 端到端流程，或移除未使用的 cron 依賴與啟動註解，讓產品行為單一且可預期。",
        "把 model、history、tool-step 與平台投遞策略集中成 typed configuration。",
        "擴充新平台時，只新增 adapter 與少量 platform capabilities；優先維持 handlers、agent 與 tools 不感知平台。",
        "若要導入更嚴格權限，新增獨立 authorization layer，根據 tenant／user／tool／side effect 決策並留下 audit event。",
    ]
    for text in roadmap:
        add_numbered(doc, text, roadmap_num_id)

    doc.add_heading("附錄：快速程式碼地圖", level=1)
    code_map = [
        ("入口", "app/api/webhooks/[platform]/route.ts"),
        ("Bot 建立", "lib/bot/index.ts"),
        ("事件處理", "lib/rina/handlers.ts"),
        ("輸入／歷史", "lib/rina/message-input.ts"),
        ("Agent／規劃", "lib/rina/agent.ts  ·  lib/rina/planner.ts"),
        ("工具註冊", "lib/rina/tools/registry.ts"),
        ("安全抓取", "lib/rina/tools/safe-fetch.ts"),
        ("輸出交付", "lib/rina/delivery.ts"),
        ("新聞零件", "lib/bot/news/rss.ts  ·  curate.ts  ·  context.ts"),
        ("本機試跑", "scripts/playground.ts"),
    ]
    for label, path in code_map:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(f"{label}　"), bold=True, color=NAVY)
        set_run_font(p.add_run(path), name="Menlo", east_asia=EAST_ASIA_FONT, size=9.2, color=INK)

    add_callout(
        doc,
        "文件基準",
        "本文件依 2026-07-25 工作區程式碼產生。若 README 與實作不一致，以程式碼行為為準；後續架構改動時，建議同步更新本文件與三張來源 SVG。",
        fill=PALE_GRAY,
        border="7B8790",
    )

    configure_sections(doc)
    return doc


def main():
    for name in ("system-context.png", "request-lifecycle.png", "tool-safety.png"):
        if not (DIAGRAM_DIR / name).exists():
            raise SystemExit(f"Missing diagram PNG: {DIAGRAM_DIR / name}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
